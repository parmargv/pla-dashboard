"""Export a financial year back into the original PLA_ACCOUNT sheet layout.

Totals, grand total, running balance and closing balance are written as Excel
formulas, so the file stays correct if someone edits a figure in Excel.
"""
from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
NUM_FMT = "#,##0.00"
DATE_FMT = "DD-MM-YYYY"
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")


def build_pla_workbook(
    fy: str,
    head_names: list[str],
    opening: dict[str, float],
    receipts: pd.DataFrame,
    expenses: pd.DataFrame,
    office_name: str = "PLA Account",
    pending: pd.DataFrame | None = None,
) -> bytes:
    """`receipts` = entries_wide(RECEIPT); `expenses` = entries_wide(EXPENSE, cleared=True)
    (the statement); `pending` = entries_wide(EXPENSE, cleared=False), listed separately."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"PLA {fy}"

    n = len(head_names)
    first_head_col = 3                       # C
    last_head_col = first_head_col + n - 1   # S for 17 heads
    total_col = last_head_col + 1            # T
    balance_col = total_col + 1              # U
    cleared_col = balance_col + 1            # V
    L = get_column_letter
    first_h, last_h, tot = L(first_head_col), L(last_head_col), L(total_col)

    def style_cell(cell, *, bold=False, fill=None, num=False, color=None, align=None):
        cell.font = Font(name=FONT, bold=bold, color=color or "000000", size=10)
        cell.border = BORDER
        if fill:
            cell.fill = fill
        if num:
            cell.number_format = NUM_FMT
        if align:
            cell.alignment = align

    def title_row(row: int, text: str):
        ws.cell(row=row, column=1, value=text)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_head_col)
        c = ws.cell(row=row, column=1)
        c.font = Font(name=FONT, bold=True, size=12, color="FFFFFF")
        c.fill = TITLE_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")

    def header_row(row: int, ref_label: str, with_balance: bool, extra: list[str] | None = None):
        headers = ["Date", ref_label, *head_names, "TOTAL"] + (["BALANCE"] if with_balance else []) + (extra or [])
        for col, text in enumerate(headers, start=1):
            c = ws.cell(row=row, column=col, value=text)
            style_cell(c, bold=True, fill=HEADER_FILL,
                       align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        ws.row_dimensions[row].height = 42

    def row_total(row: int):
        c = ws.cell(row=row, column=total_col, value=f"=SUM({first_h}{row}:{last_h}{row})")
        style_cell(c, bold=True, num=True)

    def date_cell(row: int, col: int, value):
        c = ws.cell(row=row, column=col, value=value if pd.notna(value) else None)
        style_cell(c)
        c.number_format = DATE_FMT

    def write_entries(df: pd.DataFrame, start_row: int, ref_label: str, cleared_col_no: int | None = None) -> int:
        """Writes entry rows; returns the last row used (start_row-1 if none)."""
        row = start_row - 1
        for _, rec in df.iterrows():
            row += 1
            date_cell(row, 1, rec["Date"])
            if cleared_col_no:
                date_cell(row, cleared_col_no, rec.get("Cleared On"))
            style_cell(ws.cell(row=row, column=2, value=str(rec[ref_label])))
            for i, name in enumerate(head_names):
                val = float(rec.get(name, 0) or 0)
                style_cell(ws.cell(row=row, column=first_head_col + i, value=val or None), num=True)
            row_total(row)
        return row

    # ---------------- RECEIPT ----------------
    title_row(1, f"{office_name} — RECEIPT (FY {fy})")
    header_row(2, "CHALAN NO", with_balance=False)

    opening_row = 3
    style_cell(ws.cell(row=opening_row, column=1, value="OPENING BALANCE"), bold=True, fill=TOTAL_FILL)
    ws.merge_cells(start_row=opening_row, start_column=1, end_row=opening_row, end_column=2)
    for i, name in enumerate(head_names):
        style_cell(ws.cell(row=opening_row, column=first_head_col + i, value=float(opening.get(name, 0))),
                   num=True, fill=TOTAL_FILL)
    row_total(opening_row)

    r_first = opening_row + 1
    r_last = write_entries(receipts, r_first, "Challan No")
    r_range_end = max(r_last, r_first)  # keep SUM ranges valid with zero entries

    total_receipt_row = r_range_end + 1
    grand_total_row = total_receipt_row + 1
    for label, row in (("TOTAL RECEIPT", total_receipt_row), ("GRAND TOTAL", grand_total_row)):
        style_cell(ws.cell(row=row, column=1, value=label), bold=True, fill=TOTAL_FILL)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    for col in range(first_head_col, total_col + 1):
        cl = L(col)
        style_cell(ws.cell(row=total_receipt_row, column=col, value=f"=SUM({cl}{r_first}:{cl}{r_range_end})"),
                   bold=True, num=True, fill=TOTAL_FILL)
        style_cell(ws.cell(row=grand_total_row, column=col,
                           value=f"={cl}{opening_row}+{cl}{total_receipt_row}"),
                   bold=True, num=True, fill=TOTAL_FILL)

    # ---------------- EXPENDITURE ----------------
    exp_title_row = grand_total_row + 3
    title_row(exp_title_row, f"{office_name} — EXPENDITURE (FY {fy}) — cleared cheques")
    header_row(exp_title_row + 1, "CHEQUE NO", with_balance=True, extra=["CLEARED ON"])

    e_first = exp_title_row + 2
    e_last = write_entries(expenses, e_first, "Cheque No", cleared_col_no=cleared_col)
    bal = L(balance_col)
    for row in range(e_first, e_last + 1):
        prev = f"${tot}${grand_total_row}" if row == e_first else f"{bal}{row - 1}"
        style_cell(ws.cell(row=row, column=balance_col, value=f"={prev}-{tot}{row}"), bold=True, num=True)
    e_range_end = max(e_last, e_first)

    total_expense_row = e_range_end + 1
    closing_row = total_expense_row + 1
    for label, row in (("TOTAL EXPENSE", total_expense_row), ("CLOSING BALANCE", closing_row)):
        style_cell(ws.cell(row=row, column=1, value=label), bold=True, fill=TOTAL_FILL)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    for col in range(first_head_col, total_col + 1):
        cl = L(col)
        style_cell(ws.cell(row=total_expense_row, column=col, value=f"=SUM({cl}{e_first}:{cl}{e_range_end})"),
                   bold=True, num=True, fill=TOTAL_FILL)
        style_cell(ws.cell(row=closing_row, column=col,
                           value=f"={cl}{grand_total_row}-{cl}{total_expense_row}"),
                   bold=True, num=True, fill=TOTAL_FILL, color="1F3A5F")

    # ---------------- pending cheques (not in the statement) ----------------
    if pending is not None and not pending.empty:
        p_title = closing_row + 3
        title_row(p_title, "UNCLEARED CHEQUES — issued, not yet cleared (not included above)")
        header_row(p_title + 1, "CHEQUE NO", with_balance=False)
        p_first = p_title + 2
        p_last = write_entries(pending, p_first, "Cheque No")
        p_total = p_last + 1
        style_cell(ws.cell(row=p_total, column=1, value="TOTAL PENDING"), bold=True, fill=TOTAL_FILL)
        ws.merge_cells(start_row=p_total, start_column=1, end_row=p_total, end_column=2)
        avail_row = p_total + 1
        style_cell(ws.cell(row=avail_row, column=1, value="AVAILABLE BALANCE"), bold=True, fill=TOTAL_FILL)
        ws.merge_cells(start_row=avail_row, start_column=1, end_row=avail_row, end_column=2)
        for col in range(first_head_col, total_col + 1):
            cl = L(col)
            style_cell(ws.cell(row=p_total, column=col, value=f"=SUM({cl}{p_first}:{cl}{p_last})"),
                       bold=True, num=True, fill=TOTAL_FILL)
            style_cell(ws.cell(row=avail_row, column=col, value=f"={cl}{closing_row}-{cl}{p_total}"),
                       bold=True, num=True, fill=TOTAL_FILL, color="1F3A5F")

    # ---------------- layout ----------------
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 13
    for col in range(first_head_col, cleared_col + 1):
        ws.column_dimensions[L(col)].width = 13
    ws.freeze_panes = "C3"
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
